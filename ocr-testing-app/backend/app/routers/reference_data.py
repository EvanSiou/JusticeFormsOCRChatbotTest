"""
Reference data template and document reference data management endpoints.
"""
import io
import logging
from typing import Optional
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, status
from fastapi.responses import StreamingResponse

from app.auth.dependencies import get_current_user_id, get_current_user_name
from app.services.firestore import FirestoreService
from app.services.reference_parser import parse_reference_file
from app.models.reference_data import (
    ReferenceDataTemplate,
    ReferenceField,
    CreateReferenceTemplateRequest,
    UpdateReferenceTemplateRequest,
    UpdateDocumentReferenceRequest,
)

logger = logging.getLogger(__name__)
router = APIRouter()


# ==================== Reference Template Download ====================

# Prisoner Registration Form fields organized by page/section
# These match the classification prompt field names exactly
_PRISONER_REG_FIELDS = [
    # Page 1 - Arresting Officer
    ("1", "Arresting Officer", "officer_type"),
    ("1", "Arresting Officer", "officer_last_name"),
    ("1", "Arresting Officer", "officer_first_name"),
    ("1", "Arresting Officer", "badge_number"),
    ("1", "Arresting Officer", "employee_id_number"),
    ("1", "Arresting Officer", "contact_phone_number"),
    ("1", "Arresting Officer", "agency"),
    # Page 1 - Prisoner Information
    ("1", "Prisoner Information", "first_name"),
    ("1", "Prisoner Information", "last_name"),
    ("1", "Prisoner Information", "middle_name"),
    ("1", "Prisoner Information", "ssn"),
    ("1", "Prisoner Information", "date_of_birth"),
    ("1", "Prisoner Information", "age"),
    ("1", "Prisoner Information", "sex"),
    ("1", "Prisoner Information", "race"),
    ("1", "Prisoner Information", "ethnicity"),
    ("1", "Prisoner Information", "citizenship"),
    ("1", "Prisoner Information", "country_of_birth"),
    ("1", "Prisoner Information", "city_of_birth"),
    ("1", "Prisoner Information", "height"),
    ("1", "Prisoner Information", "weight"),
    ("1", "Prisoner Information", "eyes"),
    ("1", "Prisoner Information", "skin"),
    ("1", "Prisoner Information", "hair_type"),
    ("1", "Prisoner Information", "hair_length"),
    ("1", "Prisoner Information", "hair_color"),
    ("1", "Prisoner Information", "build"),
    ("1", "Prisoner Information", "beard"),
    ("1", "Prisoner Information", "mustache"),
    ("1", "Prisoner Information", "glasses"),
    ("1", "Prisoner Information", "marital_status"),
    ("1", "Prisoner Information", "religious_preference"),
    ("1", "Prisoner Information", "veteran"),
    ("1", "Prisoner Information", "using_drugs"),
    ("1", "Prisoner Information", "using_alcohol"),
    ("1", "Prisoner Information", "wanted"),
    ("1", "Prisoner Information", "agency_wanting_person"),
    ("1", "Prisoner Information", "agency_contact_person"),
    # Page 1 - Additional Identifiers
    ("1", "Additional Identifiers", "hcso_spn"),
    ("1", "Additional Identifiers", "state_issued_id_number"),
    ("1", "Additional Identifiers", "issuing_state"),
    ("1", "Additional Identifiers", "drivers_license_number"),
    ("1", "Additional Identifiers", "dl_state"),
    ("1", "Additional Identifiers", "dl_type"),
    ("1", "Additional Identifiers", "sid_number"),
    ("1", "Additional Identifiers", "fbi_number"),
    ("1", "Additional Identifiers", "afis_number"),
    ("1", "Additional Identifiers", "so_number"),
    ("1", "Additional Identifiers", "da_log_number"),
    ("1", "Additional Identifiers", "assistant_da_name"),
    # Page 1 - Scars/Marks/Tattoos
    ("1", "Scars/Marks/Tattoos", "smt_1_type"),
    ("1", "Scars/Marks/Tattoos", "smt_1_location"),
    ("1", "Scars/Marks/Tattoos", "smt_1_description"),
    ("1", "Scars/Marks/Tattoos", "smt_2_type"),
    ("1", "Scars/Marks/Tattoos", "smt_2_location"),
    ("1", "Scars/Marks/Tattoos", "smt_2_description"),
    ("1", "Scars/Marks/Tattoos", "smt_3_type"),
    ("1", "Scars/Marks/Tattoos", "smt_3_location"),
    ("1", "Scars/Marks/Tattoos", "smt_3_description"),
    # Page 1 - Address
    ("1", "Address", "address_type"),
    ("1", "Address", "address_street"),
    ("1", "Address", "address_city"),
    ("1", "Address", "address_state"),
    ("1", "Address", "address_zip"),
    ("1", "Address", "address_source"),
    # Page 1 - Telephone
    ("1", "Telephone", "phone_1_type"),
    ("1", "Telephone", "phone_1_number"),
    ("1", "Telephone", "phone_1_source"),
    ("1", "Telephone", "phone_2_type"),
    ("1", "Telephone", "phone_2_number"),
    ("1", "Telephone", "phone_2_source"),
    # Page 1 - Emergency Contact
    ("1", "Emergency Contact", "emergency_first_name"),
    ("1", "Emergency Contact", "emergency_middle_name"),
    ("1", "Emergency Contact", "emergency_last_name"),
    ("1", "Emergency Contact", "emergency_relationship"),
    ("1", "Emergency Contact", "emergency_phone"),
    ("1", "Emergency Contact", "emergency_source"),
    # Page 1 - Employer
    ("1", "Employer", "employer_name"),
    ("1", "Employer", "occupation"),
    ("1", "Employer", "employer_address"),
    ("1", "Employer", "employer_city"),
    ("1", "Employer", "employer_state"),
    ("1", "Employer", "employer_zip"),
    ("1", "Employer", "employer_phone_type"),
    ("1", "Employer", "employer_phone"),
    ("1", "Employer", "employer_source"),
    # Page 2 - Arrest Information
    ("2", "Arrest Information", "arrest_date"),
    ("2", "Arrest Information", "arrest_time"),
    ("2", "Arrest Information", "arrest_location"),
    ("2", "Arrest Information", "arresting_agency"),
    ("2", "Arrest Information", "prisoner_health_condition"),
    ("2", "Arrest Information", "unusual_behavior"),
    # Page 2 - Arresting Officer
    ("2", "Arresting Officer (Page 2)", "arr_officer_last_name"),
    ("2", "Arresting Officer (Page 2)", "arr_officer_first_name"),
    ("2", "Arresting Officer (Page 2)", "arr_officer_badge"),
    ("2", "Arresting Officer (Page 2)", "arr_officer_employee_id"),
    ("2", "Arresting Officer (Page 2)", "arr_officer_contact"),
    ("2", "Arresting Officer (Page 2)", "arr_officer_unit"),
    ("2", "Arresting Officer (Page 2)", "arr_officer_agency"),
    # Page 2 - Transporting Officer
    ("2", "Transporting Officer", "trans_officer_last_name"),
    ("2", "Transporting Officer", "trans_officer_first_name"),
    ("2", "Transporting Officer", "trans_officer_badge"),
    ("2", "Transporting Officer", "trans_officer_employee_id"),
    ("2", "Transporting Officer", "trans_officer_contact"),
    ("2", "Transporting Officer", "trans_officer_unit"),
    ("2", "Transporting Officer", "trans_officer_agency"),
    # Page 2 - Property
    ("2", "Property", "valuable_property"),
    ("2", "Property", "bulk_property"),
    ("2", "Property", "clothing_property"),
    ("2", "Property", "other_property"),
    # Page 2 - Secure Packs
    ("2", "Secure Packs", "secure_pack_quantity"),
    ("2", "Secure Packs", "phone_numbers_opportunity"),
    ("2", "Secure Packs", "prisoner_signature"),
    ("2", "Secure Packs", "officer_signature"),
]


@router.get("/templates/download/prisoner-registration")
async def download_prisoner_registration_template():
    """Download a pre-filled Excel template for Prisoner Registration Form reference data.

    The template has 5 columns: Page, Section, Field, Raw Value, Resolved Value.
    Field names match the classification prompt exactly for direct 1:1 matching.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reference Data"

    # Header style
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    headers = ["Page", "Section", "Field", "Raw Value", "Resolved Value"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write field rows
    section_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    prev_section = None
    for row_idx, (page, section, field) in enumerate(_PRISONER_REG_FIELDS, 2):
        ws.cell(row=row_idx, column=1, value=page).border = thin_border
        cell_section = ws.cell(row=row_idx, column=2, value=section)
        cell_section.border = thin_border
        if section != prev_section:
            cell_section.fill = section_fill
        ws.cell(row=row_idx, column=3, value=field).border = thin_border
        ws.cell(row=row_idx, column=4, value="").border = thin_border  # Raw Value (user fills in)
        ws.cell(row=row_idx, column=5, value="").border = thin_border  # Resolved Value (user fills in)
        prev_section = section

    # Set column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["E"].width = 35

    # Freeze header row
    ws.freeze_panes = "A2"

    # Save to bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=prisoner_registration_reference_template.xlsx"
        },
    )


# ==================== Reference Data Templates ====================


@router.post("/templates", response_model=ReferenceDataTemplate)
async def create_reference_template(
    request: CreateReferenceTemplateRequest,
    current_user_id: str = Depends(get_current_user_id),
    current_user_name: str = Depends(get_current_user_name),
):
    """Create a new reference data template."""
    firestore = FirestoreService()
    template = await firestore.create_reference_template(
        name=request.name,
        fields=request.fields,
        created_by=current_user_id,
        created_by_name=current_user_name,
        form_type_description=request.form_type_description,
    )
    return template


@router.get("/templates", response_model=list)
async def list_reference_templates(
    current_user_id: str = Depends(get_current_user_id),
):
    """List all reference data templates."""
    firestore = FirestoreService()
    templates = await firestore.list_reference_templates()
    return templates


@router.get("/templates/{template_id}", response_model=ReferenceDataTemplate)
async def get_reference_template(
    template_id: str,
    current_user_id: str = Depends(get_current_user_id),
):
    """Get a reference data template by ID."""
    firestore = FirestoreService()
    template = await firestore.get_reference_template(template_id)
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    return template


@router.put("/templates/{template_id}", response_model=dict)
async def update_reference_template(
    template_id: str,
    request: UpdateReferenceTemplateRequest,
    current_user_id: str = Depends(get_current_user_id),
):
    """Update a reference data template."""
    firestore = FirestoreService()
    success = await firestore.update_reference_template(
        template_id=template_id,
        name=request.name,
        form_type_description=request.form_type_description,
        fields=request.fields,
    )
    if not success:
        raise HTTPException(status_code=404, detail="Template not found")
    return {"success": True}


@router.delete("/templates/{template_id}", response_model=dict)
async def delete_reference_template(
    template_id: str,
    current_user_id: str = Depends(get_current_user_id),
):
    """Delete a reference data template."""
    firestore = FirestoreService()
    success = await firestore.delete_reference_template(template_id)
    if not success:
        raise HTTPException(status_code=404, detail="Template not found")
    return {"success": True}


# ==================== Document Reference Data ====================


@router.get("/batches/{batch_id}", response_model=dict)
async def get_batch_reference_data(
    batch_id: str,
    current_user_id: str = Depends(get_current_user_id),
):
    """Get all reference data for a batch."""
    firestore = FirestoreService()
    batch = await firestore.get_batch_by_id(batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Batch not found")

    documents_ref = []
    for doc in batch.documents:
        documents_ref.append({
            "document_id": doc.id,
            "storage_path": doc.storage_path,
            "field_values": doc.field_values,
            "reference_data": doc.reference_data,
        })

    return {
        "batch_id": batch_id,
        "reference_template_id": batch.reference_template_id,
        "documents": documents_ref,
    }


@router.get("/batches/{batch_id}/documents/{document_id}", response_model=dict)
async def get_document_reference_data(
    batch_id: str,
    document_id: str,
    current_user_id: str = Depends(get_current_user_id),
):
    """Get reference data for a specific document."""
    firestore = FirestoreService()
    batch = await firestore.get_batch_by_id(batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Batch not found")

    for doc in batch.documents:
        if doc.id == document_id:
            return {
                "document_id": doc.id,
                "storage_path": doc.storage_path,
                "field_values": doc.field_values,
                "reference_data": doc.reference_data,
            }

    raise HTTPException(status_code=404, detail="Document not found in batch")


@router.put("/batches/{batch_id}/documents/{document_id}", response_model=dict)
async def update_document_reference_data(
    batch_id: str,
    document_id: str,
    request: UpdateDocumentReferenceRequest,
    current_user_id: str = Depends(get_current_user_id),
):
    """Update reference data for a specific document in a batch."""
    firestore = FirestoreService()

    # Convert to dicts for storage and build field_values for backward compat
    reference_data = [rd.model_dump() for rd in request.reference_data]
    field_values = {
        rd.field_name: rd.resolved_value or rd.raw_value
        for rd in request.reference_data
        if rd.resolved_value or rd.raw_value
    }

    success = await firestore.update_document_reference_data(
        batch_id=batch_id,
        document_id=document_id,
        reference_data=reference_data,
        field_values=field_values,
    )
    if not success:
        raise HTTPException(status_code=404, detail="Batch or document not found")
    return {"success": True}


@router.post("/batches/{batch_id}/upload", response_model=dict)
async def upload_batch_reference_data(
    batch_id: str,
    reference_file: UploadFile = File(..., description="Reference data file (Excel .xlsx or CSV)"),
    current_user_id: str = Depends(get_current_user_id),
):
    """Upload reference data from Excel/CSV and apply to ALL documents in a batch."""
    firestore = FirestoreService()

    # Validate batch exists
    batch = await firestore.get_batch_by_id(batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Batch not found")

    # Parse reference data from Excel/CSV
    ref_bytes = await reference_file.read()
    try:
        reference_data = parse_reference_file(ref_bytes, reference_file.filename)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to parse reference file: {e}",
        )

    if not reference_data:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Reference file contains no data rows",
        )

    # Build field_values for backward compat
    field_values = {
        rd["field_name"]: rd["resolved_value"] or rd["raw_value"]
        for rd in reference_data
        if rd.get("resolved_value") or rd.get("raw_value")
    }

    # Apply to ALL documents in the batch
    doc_ref = firestore.db.collection("batches").document(batch_id)
    doc = doc_ref.get()
    data = doc.to_dict()
    documents = data.get("documents", [])

    updated_count = 0
    for d in documents:
        d["reference_data"] = reference_data
        d["field_values"] = field_values
        updated_count += 1

    doc_ref.update({"documents": documents})

    logger.info(
        f"Applied {len(reference_data)} reference fields to {updated_count} documents "
        f"in batch {batch_id}"
    )

    return {
        "success": True,
        "fields_count": len(reference_data),
        "documents_updated": updated_count,
    }
