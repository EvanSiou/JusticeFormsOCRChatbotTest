"""Generate Word document with testing instructions."""
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

doc = Document()

# Title
title = doc.add_heading('OCR Testing App - Team Testing Instructions', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Quick Start
doc.add_heading('Quick Start', level=1)
doc.add_paragraph('1. Open your browser and navigate to: https://ocr-app-frontend-206256614025.us-central1.run.app')
doc.add_paragraph('2. Enter your email and password to login. If you don\'t have an account yet, contact Leo to create one for you.')
doc.add_paragraph('3. Once logged in, use the "User" dropdown filter at the top of each page to show only your own batches and test runs. This keeps everyone\'s work separate and easy to find.')

# Testing Plan
doc.add_heading('Testing Plan', level=1)
doc.add_paragraph('This plan guides you through learning the system step by step. Follow these days in order to build your understanding progressively.')

# Day 1
doc.add_heading('Day 1: Learn with Synthetic Forms', level=2)
p = doc.add_paragraph()
p.add_run('Goal: ').bold = True
p.add_run('Get comfortable with the basic workflow by using computer-generated test data. Synthetic forms have known expected values, making it easy to verify OCR accuracy.')

doc.add_heading('Step 1: Generate Synthetic Data', level=3)
doc.add_paragraph('1. Click "Synthetic Data" in the left navigation menu to open the data generation page.')
doc.add_paragraph('2. Click the blue "Generate New Batch" button in the top right corner of the page.')
doc.add_paragraph('3. In the dialog that appears, select "OP-SWO-DOCKET.PNG" from the form dropdown. This is a standard court form we\'re testing.')
doc.add_paragraph('4. Set the number of copies to 5. This creates a small batch that processes quickly while you\'re learning.')
doc.add_paragraph('5. Leave the scan quality at "Medium" for realistic degradation (noise, slight rotation, blur).')
doc.add_paragraph('6. Click "Generate" and wait 30-60 seconds. Your new batch will appear in the list below with your name attached.')

doc.add_heading('Step 2: Run Your First Test', level=3)
doc.add_paragraph('1. Click "Run Tests" in the left navigation menu to open the test execution page.')
doc.add_paragraph('2. Find your batch in the list (it will show your username). Click anywhere on the batch card to select it - a blue border indicates selection.')
doc.add_paragraph('3. In the "Layout Detection Library" dropdown, select "doclayout_yolo". This is the fastest and most accurate detector for forms.')
doc.add_paragraph('4. In the "OCR Library" dropdown, select "easyocr". This is a reliable baseline OCR engine that works well on most text.')
doc.add_paragraph('5. Click the green "Run Tests" button. A progress bar will appear showing documents processed.')
doc.add_paragraph('6. Wait for the test to complete (typically 1-2 minutes for 5 documents). The status will change to "completed" when done.')

doc.add_heading('Step 3: Review Results', level=3)
doc.add_paragraph('1. Once the test shows "completed", click the "View Results" link next to it to see the detailed results page.')
doc.add_paragraph('2. The top of the page shows overall accuracy - this is the average match score across all fields in all documents.')
doc.add_paragraph('3. Below that, you\'ll see a per-field breakdown showing which fields were recognized well and which had problems.')
doc.add_paragraph('4. Click on any document in the list to see its image with detected text regions highlighted as colored boxes.')
doc.add_paragraph('5. Compare the "Expected" values (what we filled in) versus "Extracted" values (what OCR read) to understand accuracy.')

# Day 2
doc.add_heading('Day 2: Compare OCR Libraries', level=2)
p = doc.add_paragraph()
p.add_run('Goal: ').bold = True
p.add_run('Understand how different OCR engines perform on the same documents. Running multiple tests on one batch lets you directly compare accuracy and speed.')

doc.add_paragraph('Using the SAME batch you created yesterday, run tests with each OCR engine below. Keep the layout detector as "doclayout_yolo" for all tests to ensure a fair comparison.')

# Create comparison table
table = doc.add_table(rows=6, cols=4)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Test #'
hdr_cells[1].text = 'Layout Detector'
hdr_cells[2].text = 'OCR Engine'
hdr_cells[3].text = 'What to Expect'

data = [
    ('1', 'doclayout_yolo', 'easyocr', 'Fast, reliable baseline - good starting point for comparison'),
    ('2', 'doclayout_yolo', 'paddleocr', 'Fast, often best accuracy on printed/typed text in forms'),
    ('3', 'doclayout_yolo', 'tesseract', 'Fast, classic OCR engine - works best on clean, high-contrast scans'),
    ('4', 'doclayout_yolo', 'surya', 'Medium speed, handles mixed layouts and rotated text well'),
    ('5', 'doclayout_yolo', 'trocr', 'Slower, specialized for handwritten text - try this last'),
]
for i, row_data in enumerate(data):
    row = table.rows[i + 1].cells
    for j, text in enumerate(row_data):
        row[j].text = text

doc.add_paragraph('')
doc.add_paragraph('After running all 5 tests, go to the Results page and compare the accuracy percentages. Note which OCR engine performed best on your specific form - results may vary by form type.')

# Day 3
doc.add_heading('Day 3: Test VLM Engines (Vision Language Models)', level=2)
p = doc.add_paragraph()
p.add_run('Goal: ').bold = True
p.add_run('Test the most advanced AI-powered OCR engines. These are large neural networks that understand documents holistically rather than detecting regions first.')

p = doc.add_paragraph()
p.add_run('IMPORTANT DIFFERENCE: ').bold = True
p.add_run('VLM engines (got_ocr and mineru) work differently from traditional OCR:')
doc.add_paragraph('• They process the ENTIRE PAGE in a single pass, like a human reading a document', style='List Bullet')
doc.add_paragraph('• They do NOT need a separate layout detection step - the layout dropdown will automatically hide when you select them', style='List Bullet')
doc.add_paragraph('• They are significantly slower (2-3 minutes per document vs seconds) because they\'re running large AI models', style='List Bullet')
doc.add_paragraph('• They can potentially handle complex layouts, tables, and mixed content better than traditional OCR', style='List Bullet')

table2 = doc.add_table(rows=3, cols=2)
table2.style = 'Table Grid'
hdr_cells = table2.rows[0].cells
hdr_cells[0].text = 'OCR Engine'
hdr_cells[1].text = 'Description'
table2.rows[1].cells[0].text = 'got_ocr'
table2.rows[1].cells[1].text = 'GOT-OCR2.0 by StepFun AI - 580 million parameter vision-language model designed for general OCR tasks'
table2.rows[2].cells[0].text = 'mineru'
table2.rows[2].cells[1].text = 'MinerU 2.5 by OpenDataLab - 1.2 billion parameter model based on Qwen2VL, excellent for document parsing'

doc.add_paragraph('')
doc.add_paragraph('To test VLM engines:')
doc.add_paragraph('1. Go to Run Tests and select your batch as usual.')
doc.add_paragraph('2. In the OCR dropdown, select "got_ocr" or "mineru".')
doc.add_paragraph('3. Notice the Layout dropdown disappears and a blue message explains why.')
doc.add_paragraph('4. Click Run Tests. Be patient - expect 2-3 minutes per document.')
doc.add_paragraph('5. Compare the results to your Day 2 tests to see if the VLM performed better.')

# Day 4
doc.add_heading('Day 4: Handwritten Forms', level=2)
p = doc.add_paragraph()
p.add_run('Goal: ').bold = True
p.add_run('Test OCR on actual handwritten documents, which are much harder to read than printed text. This simulates real-world court forms filled in by hand.')

doc.add_heading('Upload Handwritten Batch', level=3)
doc.add_paragraph('1. First, scan or photograph some handwritten court forms (or use provided samples).')
doc.add_paragraph('2. Go to the Synthetic Data page and click "Upload Handwritten Batch" instead of Generate.')
doc.add_paragraph('3. Select your scanned image files (PNG, JPG, or PDF format supported).')
doc.add_paragraph('4. The system will create distortion copies (rotated, noisy versions) to simulate real scanning conditions.')
doc.add_paragraph('5. Your handwritten batch will appear with a purple "Handwritten" badge to distinguish it from synthetic batches.')

doc.add_heading('Run Handwritten Tests', level=3)
doc.add_paragraph('1. Go to Run Tests and select your handwritten batch (look for the purple "Handwritten" badge).')
doc.add_paragraph('2. Notice that Layout Detection is automatically disabled for handwritten batches - these use full-page OCR only.')
doc.add_paragraph('3. For handwritten text, try these OCR engines in order of recommendation:')
doc.add_paragraph('   • trocr - Microsoft\'s transformer model specifically trained on handwritten text', style='List Bullet')
doc.add_paragraph('   • got_ocr - VLM that can handle mixed printed and handwritten content', style='List Bullet')
doc.add_paragraph('   • easyocr - General purpose, sometimes works on clear handwriting', style='List Bullet')
doc.add_paragraph('4. Run tests and compare which engine reads your handwriting most accurately.')

doc.add_heading('Verify Handwritten Results', level=3)
doc.add_paragraph('1. Go to the Verify page in the left navigation.')
doc.add_paragraph('2. Find your handwritten test run and click "Verify" to open the verification interface.')
doc.add_paragraph('3. For each document, mark which detected text regions contain actual handwriting vs printed template text.')
doc.add_paragraph('4. This "important" tagging helps calculate meaningful accuracy - we only want to measure handwritten portions.')
doc.add_paragraph('5. Correct any misread text by typing the actual value.')
doc.add_paragraph('6. Submit your verifications to update the accuracy metrics.')

# Library Reference
doc.add_heading('Library Reference', level=1)
doc.add_paragraph('Use this reference to understand what each library does and when to use it.')

doc.add_heading('Layout Detectors', level=2)
doc.add_paragraph('Layout detectors find text regions on the page before OCR reads them. Choose based on your document type:')
table3 = doc.add_table(rows=4, cols=3)
table3.style = 'Table Grid'
hdr = table3.rows[0].cells
hdr[0].text = 'Detector'
hdr[1].text = 'Speed'
hdr[2].text = 'Best Used For'
data3 = [
    ('doclayout_yolo', 'Fast (1-2 sec)', 'Standard forms, documents with clear structure - RECOMMENDED for most cases'),
    ('doctr', 'Medium (3-5 sec)', 'Dense documents with lots of text blocks, multi-column layouts'),
    ('surya', 'Medium (3-5 sec)', 'Mixed layouts with tables, figures, and text in various orientations'),
]
for i, d in enumerate(data3):
    row = table3.rows[i + 1].cells
    for j, t in enumerate(d):
        row[j].text = t

doc.add_paragraph('')
doc.add_heading('OCR Engines', level=2)
doc.add_paragraph('OCR engines read text from detected regions. Speed and accuracy vary significantly:')
table4 = doc.add_table(rows=8, cols=4)
table4.style = 'Table Grid'
hdr = table4.rows[0].cells
hdr[0].text = 'Engine'
hdr[1].text = 'Speed'
hdr[2].text = 'Best Used For'
hdr[3].text = 'Additional Notes'
data4 = [
    ('easyocr', 'Fast (seconds)', 'General printed text', 'Good baseline, works on most documents'),
    ('paddleocr', 'Fast (seconds)', 'Printed forms and documents', 'Often achieves highest accuracy on clean printed text'),
    ('tesseract', 'Fast (seconds)', 'High-contrast clean scans', 'Classic open-source OCR, requires good image quality'),
    ('surya', 'Medium (5-10 sec)', 'Mixed layouts', 'Handles rotated text and complex layouts well'),
    ('trocr', 'Slow (30-60 sec)', 'Handwritten text', 'Microsoft\'s specialized handwriting model - BEST for handwriting'),
    ('got_ocr', 'Very Slow (2-3 min)', 'Full documents', 'VLM - processes entire page, no layout detection needed'),
    ('mineru', 'Very Slow (2-3 min)', 'Full documents', 'VLM - largest model, best for complex document understanding'),
]
for i, d in enumerate(data4):
    row = table4.rows[i + 1].cells
    for j, t in enumerate(d):
        row[j].text = t

doc.add_paragraph('')
doc.add_heading('Scan Quality Settings', level=2)
doc.add_paragraph('When generating synthetic data, you can choose how much degradation to apply:')
doc.add_paragraph('• Light: Minimal noise, very slight rotation (±1°) - simulates high-quality scanner', style='List Bullet')
doc.add_paragraph('• Medium: Moderate noise and blur, some rotation (±3°) - simulates typical office scanner', style='List Bullet')
doc.add_paragraph('• Heavy: Significant noise, blur, and rotation (±5°) - simulates poor quality scans or phone photos', style='List Bullet')

# Tips
doc.add_heading('Tips for Effective Testing', level=1)
doc.add_paragraph('1. ALWAYS use the User filter: This keeps your batches and tests separate from other team members. Set it once when you log in.')
doc.add_paragraph('2. Start with small batches: Use 5-10 documents while learning. You can generate larger batches (50-100) once you\'re comfortable.')
doc.add_paragraph('3. Compare apples to apples: Run multiple OCR engines on the SAME batch to fairly compare their accuracy.')
doc.add_paragraph('4. Be patient with VLMs: got_ocr and mineru take 2-3 minutes per document. Don\'t assume it\'s stuck - check the progress.')
doc.add_paragraph('5. Use Verify to improve accuracy: The verification page lets you correct OCR mistakes, which improves our ground truth data.')
doc.add_paragraph('6. Document your findings: Note which engine worked best for which form type - this helps the whole team.')

# Troubleshooting
doc.add_heading('Troubleshooting Common Issues', level=1)

p = doc.add_paragraph()
p.add_run('Test stuck at 0% with no progress').bold = True
doc.add_paragraph('• The test may have failed silently. Look for any red error messages on the page.')
doc.add_paragraph('• Click "Cancel" to stop the stuck test, then try running it again.')
doc.add_paragraph('• If it fails repeatedly, try a different OCR engine or contact Leo.')

p = doc.add_paragraph()
p.add_run('VLM tests (got_ocr/mineru) are extremely slow').bold = True
doc.add_paragraph('• This is expected - these are 580M-1.2B parameter AI models running on CPU.')
doc.add_paragraph('• Each document takes 2-3 minutes. A 10-document batch takes 20-30 minutes.')
doc.add_paragraph('• For quick testing, use easyocr or paddleocr instead.')
doc.add_paragraph('• GPU acceleration is being set up to speed this up in the future.')

p = doc.add_paragraph()
p.add_run('"No batches available" message').bold = True
doc.add_paragraph('• You need to generate synthetic data first before running tests.')
doc.add_paragraph('• Check the User filter dropdown - it might be filtering to a different user.')
doc.add_paragraph('• Try setting User filter to "All Users" to see all batches.')

p = doc.add_paragraph()
p.add_run('Can\'t find my test results').bold = True
doc.add_paragraph('• Check the User filter on the Run Tests page - make sure it shows your username.')
doc.add_paragraph('• The test might still be running - look for a progress bar at the top of the page.')
doc.add_paragraph('• Tests are listed most recent first - scroll down if you ran it a while ago.')

p = doc.add_paragraph()
p.add_run('Poor accuracy results').bold = True
doc.add_paragraph('• Try a different OCR engine - some work better on certain form types.')
doc.add_paragraph('• Check the scan quality setting - "Heavy" degradation makes OCR harder.')
doc.add_paragraph('• For handwritten text, use trocr or got_ocr instead of general OCR engines.')

# Contact
doc.add_heading('Questions or Problems?', level=1)
doc.add_paragraph('Contact Leo for assistance with:')
doc.add_paragraph('• Account creation - need login credentials for new team members', style='List Bullet')
doc.add_paragraph('• Technical issues - errors, crashes, or unexpected behavior', style='List Bullet')
doc.add_paragraph('• Feature requests - ideas for improving the testing workflow', style='List Bullet')
doc.add_paragraph('• Adding new forms - need to upload new court form templates', style='List Bullet')

# Save
doc.save('OCR_Testing_Instructions.docx')
print('Created: OCR_Testing_Instructions.docx')
